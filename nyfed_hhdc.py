"""
NY Fed Quarterly Report on Household Debt and Credit (HHDC) -> SQLite DB (output/db/nyfed_household_credit.db + .gz)
FRED에는 없는 데이터라 NY Fed가 분기마다 올리는 엑셀(HHD_C_Report_YYYYQN.xlsx)을 받아 읽는다.

    python nyfed_hhdc.py              # 최신 분기 보고서를 찾아서 반영
    python nyfed_hhdc.py --quarter 2026Q2

발표: 분기 종료 약 6주 뒤 (2·5·8·11월). 아직 안 나온 분기 주소는 엑셀 대신 HTML 페이지가 와서 건너뜀.

구조 (build_db.py가 만드는 FRED 항목별 DB와 같음 -> 대시보드가 같은 방식으로 읽음)
    series_info   : series_id = P<쪽>_<항목> (예: P12_CC), title = '<그림 제목>: <항목>', units = 시트 2행
    obs_quarterly : (series_id, date, value), date = 분기 첫날 (03:Q1 -> 2003-01-01)
엑셀의 'Page N Data' 시트 중 전국 분기 시계열(3~30쪽)만 읽고, 주별·한 시점 단면 시트(21, 32쪽~)는 뺀다.
보고서마다 과거값이 수정될 수 있어 매번 전체를 다시 쓴다.
"""
import io
import re
import sqlite3
import warnings
from argparse import ArgumentParser
from datetime import date, datetime

import openpyxl
import requests
from urllib3.exceptions import InsecureRequestWarning

from config import user_agent
from util import make_dir, write_gz

warnings.filterwarnings("ignore", category=InsecureRequestWarning)

URL = 'https://www.newyorkfed.org/medialibrary/interactives/householdcredit/data/xls/HHD_C_Report_{q}.xlsx'
DB_PATH = './output/db/nyfed_household_credit.db'
PAGES = list(range(3, 21)) + list(range(22, 31))  # 전국 분기 시계열 시트
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def recent_quarters(n=6):
    today = date.today()
    y, q = today.year, (today.month - 1) // 3 + 1
    for _ in range(n):
        yield f'{y}Q{q}'
        y, q = (y, q - 1) if q > 1 else (y - 1, 4)


def download(quarter=None):
    for q in [quarter] if quarter else recent_quarters():
        r = requests.get(URL.format(q=q), headers={'User-Agent': user_agent}, timeout=60, verify=False)
        if r.status_code == 200 and r.headers.get('Content-Type', '').startswith(XLSX):
            return q, r.content
        print(f'    {q}: 아직 없음')
    raise SystemExit('받을 수 있는 보고서가 없어요.')


def quarter_start(a):
    if isinstance(a, datetime):  # 분기 말월(3·6·9·12월 1일)로 들어온 경우
        return date(a.year, a.month - 2, 1)
    m = re.fullmatch(r'(\d\d):Q([1-4])', str(a).strip())
    if not m:
        return None
    yy = int(m.group(1))
    return date(1900 + yy if yy >= 90 else 2000 + yy, (int(m.group(2)) - 1) * 3 + 1, 1)


def slug(text):
    return re.sub(r'[^A-Z0-9]+', '_', str(text).upper()).strip('_')


def parse_sheet(ws, page):
    rows = list(ws.iter_rows(values_only=True))
    title = str(rows[0][0]).strip()
    units = str(rows[1][0]).strip() if rows[1][0] else ''
    if units.startswith('Return to') or units.startswith('* Credit'):  # 신용점수 분포 시트는 단위 행이 없음
        units = 'Score'
    first = next(i for i, r in enumerate(rows) if r and quarter_start(r[0]))
    header = rows[first - 1]
    cols = []
    for j in range(1, len(header)):  # A열(분기) 다음부터 이어진 헤더만 (옆에 붙은 다른 표·메모는 무시)
        h = header[j]
        if not isinstance(h, str) or not h.strip():
            break
        cols.append((j, h.strip()))
    series = {f'P{page:02d}_{slug(h)}': (f'{title}: {h}', units, {}) for _, h in cols}
    current = None
    for r in rows[first:]:
        if not r:
            continue
        d = quarter_start(r[0])
        if d:
            current = d
        elif r[0] is not None:  # 표가 끝나고 메모 등
            break
        if current is None:
            continue
        for j, h in cols:  # 10쪽처럼 같은 분기 값이 다음 줄(A열 빈칸)에 이어지는 경우도 current로 받음
            v = r[j] if j < len(r) else None
            if isinstance(v, (int, float)):
                series[f'P{page:02d}_{slug(h)}'][2][current.isoformat()] = float(v)
    return {k: v for k, v in series.items() if v[2]}


def main():
    parser = ArgumentParser(description='NY Fed 가계신용 보고서 -> sqlite')
    parser.add_argument('--quarter', help='예: 2026Q2 (기본: 최신)')
    args = parser.parse_args()

    q, content = download(args.quarter)
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    all_series = {}
    for page in PAGES:
        name = f'Page {page} Data'
        if name in wb.sheetnames:
            all_series.update(parse_sheet(wb[name], page))

    make_dir('./output/db')
    con = sqlite3.connect(DB_PATH)
    con.execute('DROP TABLE IF EXISTS series_info')
    con.execute('DROP TABLE IF EXISTS obs_quarterly')
    con.execute('''CREATE TABLE series_info (
        series_id TEXT PRIMARY KEY, list_name TEXT, group_path TEXT, title TEXT,
        frequency TEXT, units TEXT, seasonal_adjustment TEXT,
        observation_start TEXT, observation_end TEXT, last_updated TEXT,
        n_obs INTEGER, fetched_at TEXT, status TEXT, error TEXT)''')
    con.execute('CREATE TABLE obs_quarterly (series_id TEXT, date TEXT, value REAL, PRIMARY KEY (series_id, date)) WITHOUT ROWID')
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    for sid, (title, units, obs) in all_series.items():
        days = sorted(obs)
        con.executemany('INSERT INTO obs_quarterly VALUES (?,?,?)', [(sid, d, obs[d]) for d in days])
        con.execute('INSERT INTO series_info VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)', (
            sid, 'ids_Q', title.split(':')[0], title, 'Q', units, 'NSA', days[0], days[-1],
            f'HHD_C_Report_{q}', len(days), now, 'ok', None))
    con.commit()
    con.execute('VACUUM')
    con.close()
    print(f'>>> HHD_C_Report_{q}: 시리즈 {len(all_series)}개 -> {DB_PATH} (압축본 {write_gz(DB_PATH) / 1e6:.2f}MB)')


if __name__ == '__main__':
    main()
